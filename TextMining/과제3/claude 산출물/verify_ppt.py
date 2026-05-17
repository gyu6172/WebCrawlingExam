from pptx import Presentation
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

p = Presentation('과제_3_Word2Vec_구현_학번22011831_이름김형규_완성.pptx')
for i, slide in enumerate(p.slides):
    print(f'\n=== Slide {i+1} ===')
    for shp in slide.shapes:
        if shp.has_text_frame:
            tx = shp.text_frame.text.strip()
            if tx:
                first_line = tx.split('\n')[0][:80]
                line_count = len(tx.split('\n'))
                print(f'  [text {line_count}lines] {first_line}')
        else:
            print(f'  [pic/shape] {shp.name}')
