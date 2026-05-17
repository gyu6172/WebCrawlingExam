from pptx import Presentation
from pptx.util import Emu
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

p = Presentation('과제_3_Word2Vec_구현_학번22011831_이름김형규.pptx')
print(f'Slide width: {p.slide_width}, height: {p.slide_height}')
print(f'Total slides: {len(p.slides)}')
for i, slide in enumerate(p.slides):
    print(f'\n=== Slide {i+1} (layout: {slide.slide_layout.name}) ===')
    for j, shape in enumerate(slide.shapes):
        print(f'  Shape {j}: type={shape.shape_type}, name={shape.name}')
        print(f'    pos=({shape.left}, {shape.top}), size=({shape.width}, {shape.height})')
        if shape.has_text_frame:
            for k, para in enumerate(shape.text_frame.paragraphs):
                for run in para.runs:
                    print(f'    para{k}: "{run.text}" font={run.font.name} size={run.font.size} bold={run.font.bold}')
