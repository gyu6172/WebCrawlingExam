"""기존 템플릿을 복사한 뒤 각 슬라이드에 코드/결과/이유/분석을 채워넣는다.
- Slide 1 (표지): 학생 정보 갱신
- Slide 2 (공지): 그대로 둠 (제출 시 학생이 직접 삭제)
- Slide 3~9: Step 1~7 (총 7개)
"""
import os
import shutil
import copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUT = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(OUT, "과제_3_Word2Vec_구현_학번22011831_이름김형규.pptx")
DST = os.path.join(OUT, "과제_3_Word2Vec_구현_학번22011831_이름김형규_완성.pptx")
shutil.copy(SRC, DST)

prs = Presentation(DST)
SW, SH = prs.slide_width, prs.slide_height  # 12192000 x 6858000

# === 1. 표지 슬라이드(인덱스 0) 학생 정보 갱신 ===
slide_cover = prs.slides[0]
for shape in slide_cover.shapes:
    if not shape.has_text_frame:
        continue
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            t = run.text
            if t == "00000 ":
                run.text = "데이터사이언스 "
            elif t == ": 000000":
                run.text = ": 22011831"
            elif t == ":  000":
                run.text = ":  김형규"

# === 2. 템플릿에 있는 Step 1, 2, 3 슬라이드를 4개 더 복제해 Step 4, 5, 6, 7 만들기 ===
# 현재 슬라이드 인덱스: 0=표지, 1=공지, 2=Step1, 3=Step2, 4=Step3
# 4개 복제하여 인덱스 5,6,7,8 = Step4, Step5, Step6, Step7

def duplicate_slide(prs, src_idx):
    src = prs.slides[src_idx]
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        sp = shp._element
        sp.getparent().remove(sp)
    for shp in src.shapes:
        new.shapes._spTree.append(copy.deepcopy(shp._element))
    for rel in src.part.rels.values():
        if "notesSlide" in rel.reltype:
            continue
        new.part.rels.get_or_add(rel.reltype, rel.target_part)
    return new

# Step 3 슬라이드(인덱스 4)를 4번 복제 → 새 슬라이드 4개 추가
for _ in range(4):
    duplicate_slide(prs, 4)

# === 3. 각 Step에 들어갈 컨텐츠 ===
STEP_DATA = [
    {
        "step": "Step 1",
        "code_img": "code_step1.png",
        "out_img": "out_step1.png",
        "reasons": [
            "Word2Vec 학습은 단어를 정수 인덱스로 매핑한 사전이 필수이므로, 코퍼스를 토큰화하고 set()으로 중복을 제거해 word_dict / number_dict를 만들었다.",
            "이후 단계의 원-핫 벡터 차원과 출력층 크기가 모두 vocab_size에 의존하므로, 사전 구축이 가장 먼저 수행되어야 한다.",
        ],
        "analysis": [
            "총 어휘 수는 15개(vocab_size=15)로 매우 작은 toy 코퍼스이다. 작은 어휘는 학습 시간을 짧게 해주지만, 의미 관계가 형성되려면 동일 패턴이 여러 번 등장해야 한다.",
            "코퍼스 내에 'king/queen', 'man/woman', 'seoul/tokyo/paris' 같은 의미적 짝이 의도적으로 배치되어 있어, Skip-gram이 잘 학습되면 시각화에서 군집이 형성될 것이라 기대된다.",
            "set() 자료형은 순서가 보장되지 않아 word_dict 인덱스가 실행마다 달라질 수 있는데, 이는 학습 결과 시각화 좌표에는 영향이 없고 단순히 인덱스 매핑만 바뀐다.",
        ],
    },
    {
        "step": "Step 2",
        "code_img": "code_step2.png",
        "out_img": "out_step2.png",
        "reasons": [
            "Skip-gram은 중심 단어로 주변 단어를 예측하므로 (target, context) 쌍 자체가 학습 데이터가 된다. window_size 내의 인덱스를 순회해 자기 자신을 제외한 모든 주변 단어를 짝지어 주었다.",
            "tokens[j]를 word_dict로 매핑해 인덱스 형태로 저장한 이유는 CrossEntropyLoss가 정답 라벨로 정수 인덱스를 요구하기 때문이다.",
        ],
        "analysis": [
            "window_size=1이라는 좁은 윈도우에서도 7개 문장만으로 42개의 학습 쌍이 생성됐다. 같은 문장 안에서 (target, context)와 (context, target)이 모두 들어가 양방향 학습이 자연스럽게 이뤄진다.",
            "예: 'he is a king' → (he,is)(is,he)(is,a)(a,is)(a,king)(king,a). 'is'와 'a'가 거의 모든 문장에 등장해 학습 쌍에서 가장 자주 나타나며, 이는 일반적 의미를 갖는 단어가 어휘 공간 중심부에 위치하는 결과로 이어진다.",
            "윈도우가 작으면 직접 인접 단어만 학습하므로 의미 군집은 명확해지지만, 멀리 떨어진 단어 간 관계(예: 같은 문장의 첫 단어와 마지막 단어)는 직접적으로 학습되지 않는다.",
        ],
    },
    {
        "step": "Step 3",
        "code_img": "code_step3.png",
        "out_img": None,
        "reasons": [
            "Skip-gram은 본질적으로 V→D→V의 두 단 선형 변환이므로, bias 없는 두 개의 nn.Linear만으로 구현 가능하다. W는 입력층-은닉층(임베딩) 행렬, WT는 은닉층-출력층 행렬이다.",
            "임베딩 시각화를 2차원 평면에서 직접 확인하려고 embedding_size=2로 설정했다. 실 사용에서는 보통 100~300차원을 쓴다.",
        ],
        "analysis": [
            "은닉층에 비선형 활성화가 없는 점이 핵심이다. 이는 Word2Vec이 행렬 분해(matrix factorization) 관점으로 해석될 수 있는 이유와 직결된다.",
            "최종 임베딩으로 사용하는 것은 W의 가중치 행렬이며, WT는 학습 보조용이다. 모델 파라미터 수는 V*D + D*V = 15*2 + 2*15 = 60개로 매우 작다.",
            "bias=False로 두는 이유는 원-핫 입력 + 단순 lookup 구조에서 bias가 단어별 임베딩 의미를 흐리기 때문이다.",
        ],
    },
    {
        "step": "Step 4",
        "code_img": "code_step4.png",
        "out_img": "out_step4.png",
        "reasons": [
            "출력은 어휘 전체에 대한 점수이므로 다중 클래스 분류로 보고 CrossEntropyLoss를 선택했다. 이 손실은 내부에 softmax가 포함되어 별도 softmax 호출이 필요 없다.",
            "Adam 옵티마이저는 작은 코퍼스에서도 학습률을 자동 조절해 빠르게 수렴하므로 채택했다.",
            "get_one_hot()으로 단어 인덱스를 V차원 원-핫 벡터로 변환해야 nn.Linear의 입력 형태와 맞는다.",
        ],
        "analysis": [
            "원-핫 벡터를 W에 통과시키는 연산은 사실상 W의 해당 단어 행을 그대로 가져오는 'lookup'과 동일하다. 행렬곱처럼 보이지만 실제로는 임베딩 테이블 조회이다.",
            "실무에서는 이 비효율을 피하려고 nn.Embedding을 쓰지만, 본 과제는 원-핫 → Linear의 흐름을 명시적으로 보여주기 위한 교육적 구현이다.",
            "출력에서 인덱스 0 위치만 1.0인 길이 15짜리 벡터가 잘 만들어졌음을 확인했다.",
        ],
    },
    {
        "step": "Step 5",
        "code_img": "code_step5.png",
        "out_img": "fig_loss_base.png",
        "reasons": [
            "각 (target, context) 쌍마다 forward → loss → backward → step 흐름을 반복해 W, WT 가중치를 업데이트한다. zero_grad()로 이전 gradient를 초기화하지 않으면 누적되어 학습이 망가지므로 매 스텝에서 호출했다.",
            "epoch마다 평균 loss를 기록해 학습 진행 상황을 확인하고 시각화하기 쉽도록 했다.",
        ],
        "analysis": [
            "초기 loss는 약 2.71에서 시작해 1000 epoch 만에 1.14 부근으로 빠르게 떨어진 뒤 거의 평탄해진다. 모델이 toy 코퍼스의 가능한 (target, context) 분포를 거의 다 학습했음을 의미한다.",
            "loss가 0이 아니라 1.14 근처에서 멈추는 이유는, 같은 target이 다수의 다른 context를 가지므로(예: 'is' → he/she/a/korea/japan…) 한 번의 예측으로 모든 정답을 동시에 맞출 수 없기 때문이다. 이는 본질적인 cross-entropy 하한이다.",
            "saddle 없이 매끄럽게 수렴하는 모습은 Adam의 적응형 학습률 덕분이며, SGD를 사용하면 더 흔들리고 수렴 속도도 느려질 가능성이 크다.",
        ],
    },
    {
        "step": "Step 6",
        "code_img": "code_step6.png",
        "out_img": "fig_embedding_base.png",
        "reasons": [
            "학습된 W의 행 벡터가 곧 단어의 임베딩이므로 model.W.weight.data를 numpy 배열로 변환했다. .T를 취해 (V, D) 형태로 만든 뒤 단어별 좌표로 사용했다.",
            "embedding_size=2이므로 별도 차원 축소(PCA, t-SNE) 없이 그대로 산점도로 시각화 가능하다.",
        ],
        "analysis": [
            "결과 그래프에서 의미적으로 비슷한 단어가 같은 영역에 군집을 이루었다. 좌측 하단에 king/queen/man/woman 같은 사람·성별 관련 단어가 모였고, 우측 상단에 도시명(seoul/tokyo/paris)과 대명사(he/she)가, 우측 중앙에 국가명(korea/japan/france)이 모여 있다.",
            "조사·관사 역할의 'is', 'a'는 원점 근처에 자리잡고 있는데, 거의 모든 문장에 등장해 특정 카테고리에 치우치지 않은 결과로 해석된다.",
            "이러한 군집화는 Skip-gram이 분포 가설(distributional hypothesis: 비슷한 문맥에서 등장하는 단어는 비슷한 의미를 가진다)을 잘 반영했다는 직접적 증거이다.",
        ],
    },
    {
        "step": "Step 7",
        "code_img": "code_step7.png",
        "out_img": "fig_loss_compare.png",
        "reasons": [
            "embedding_size를 2→10으로 늘려 표현력을 확장하고, window_size를 1→2로 늘려 더 넓은 문맥을 학습 신호로 사용했다. 두 변화가 학습 양상에 미치는 영향을 함께 관찰하려는 의도이다.",
            "epoch는 5000→3000으로 줄여 학습 비용 변화도 같이 점검했다.",
        ],
        "analysis": [
            "window_size=2로 확장하니 학습 쌍이 42개 → 70개로 약 +67% 늘었다. 한 target 단어가 예측해야 하는 context의 종류도 많아져 per-pair cross-entropy loss는 base(1.14)보다 오히려 높은 1.55 부근에서 수렴했다.",
            "loss 값이 더 크다고 해서 모델 품질이 나쁜 것은 아니다. 윈도우가 커질수록 같은 target이 더 다양한 context를 예측해야 하므로 이론적 cross-entropy 하한 자체가 올라간다. 즉 loss 절댓값이 아니라 학습 곡선의 안정성과 임베딩 품질로 평가해야 한다.",
            "embedding_size 2→10 증가로 의미 표현 공간이 더 풍부해져 단어 간 미세한 차이까지 표현 가능해진다는 trade-off가 있다. 다만 시각화는 직접 불가능해지므로 PCA 등 차원 축소가 필요하다.",
            "실 응용에서는 시각화보다 임베딩 품질이 중요하므로 embedding_size 100~300이 일반적이며, window_size도 코퍼스 특성과 도메인에 따라 조절한다.",
        ],
    },
]

assert len(STEP_DATA) == 7

def fill_step_slide(slide, data):
    body_box, step_box = None, None
    for shp in slide.shapes:
        if not shp.has_text_frame:
            continue
        tx = shp.text_frame.text
        if tx.startswith("Step "):
            step_box = shp
        elif "현 PPT" in tx or "PPT 양식" in tx:
            body_box = shp

    if step_box is not None:
        tf = step_box.text_frame
        for i, para in enumerate(tf.paragraphs):
            for j, run in enumerate(para.runs):
                if i == 0 and j == 0:
                    run.text = data["step"]
                else:
                    run.text = ""

    if body_box is not None:
        sp = body_box._element
        sp.getparent().remove(sp)

    LEFT = 200000
    TOP_CODE = 1300000
    CODE_W = 5800000  # 좌측 영역 너비 약 5.8M EMU

    # 각 Step별 코드/출력 이미지 높이 직접 지정 (슬라이드 범위 보장)
    SLIDE_BOTTOM_LIMIT = 6500000  # 슬라이드 높이 6.858M, 여유 둠

    code_path = os.path.join(OUT, data["code_img"])
    out_path = os.path.join(OUT, data["out_img"]) if data.get("out_img") else None

    step = data["step"]
    if step == "Step 6":
        # 코드는 작게(상단), 임베딩 그래프는 메인으로 크게
        code_pic = slide.shapes.add_picture(code_path, LEFT, TOP_CODE, width=4400000)
        out_top = TOP_CODE + code_pic.height + 50000
        slide.shapes.add_picture(out_path, LEFT, out_top, height=SLIDE_BOTTOM_LIMIT - out_top)
    elif step == "Step 7":
        code_pic = slide.shapes.add_picture(code_path, LEFT, TOP_CODE, width=CODE_W)
        out_top = TOP_CODE + code_pic.height + 80000
        slide.shapes.add_picture(out_path, LEFT, out_top, height=SLIDE_BOTTOM_LIMIT - out_top)
    elif step == "Step 5":
        code_pic = slide.shapes.add_picture(code_path, LEFT, TOP_CODE, width=CODE_W)
        out_top = TOP_CODE + code_pic.height + 80000
        slide.shapes.add_picture(out_path, LEFT, out_top, height=SLIDE_BOTTOM_LIMIT - out_top)
    elif step == "Step 3":
        slide.shapes.add_picture(code_path, LEFT, TOP_CODE, width=CODE_W)
    else:
        code_pic = slide.shapes.add_picture(code_path, LEFT, TOP_CODE, width=CODE_W)
        if out_path:
            out_top = TOP_CODE + code_pic.height + 100000
            # 출력 이미지는 작은 텍스트 박스 — 원본 비율 유지
            slide.shapes.add_picture(out_path, LEFT, out_top, width=CODE_W)

    RIGHT_LEFT = 6200000
    RIGHT_TOP = 1300000
    RIGHT_W = SW - RIGHT_LEFT - 200000
    RIGHT_H = SH - RIGHT_TOP - 300000

    txbox = slide.shapes.add_textbox(RIGHT_LEFT, RIGHT_TOP, RIGHT_W, RIGHT_H)
    tf = txbox.text_frame
    tf.word_wrap = True

    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.LEFT
    r1 = p1.add_run()
    r1.text = "코드 작성 이유"
    r1.font.bold = True
    r1.font.size = Pt(20)
    r1.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    r1.font.name = "맑은 고딕"

    for i, reason in enumerate(data["reasons"], 1):
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = f"{i}. {reason}"
        r.font.size = Pt(13)
        r.font.name = "맑은 고딕"
        p.space_after = Pt(4)

    tf.add_paragraph().add_run().text = ""

    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "분석 결과"
    r2.font.bold = True
    r2.font.size = Pt(20)
    r2.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    r2.font.name = "맑은 고딕"

    for i, ana in enumerate(data["analysis"], 1):
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = f"{i}. {ana}"
        r.font.size = Pt(13)
        r.font.name = "맑은 고딕"
        p.space_after = Pt(4)


# 슬라이드 인덱스: 0=표지, 1=공지, 2~8 = Step1~7
for i, data in enumerate(STEP_DATA):
    slide = prs.slides[i + 2]
    fill_step_slide(slide, data)

prs.save(DST)
print(f"Saved: {DST}")
print(f"Total slides: {len(prs.slides)}")
