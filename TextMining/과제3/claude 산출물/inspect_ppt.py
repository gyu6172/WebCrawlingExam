from pptx import Presentation
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

p = Presentation('과제_3_Word2Vec_구현_학번22011831_이름김형규.pptx')
for i, slide in enumerate(p.slides):
    print(f'--- Slide {i+1} ---')
    for shape in slide.shapes:
        if shape.has_text_frame:
            print(shape.text_frame.text)
        print()
